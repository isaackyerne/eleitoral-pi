"""Correspondência de seções entre 2022 e 2026 e Lula por GRUPO de seções.

A seção do TSE não é um código estável de ponta a ponta:
  - uma seção AGREGADA não aparece no arquivo de votos — seus votos entram na seção
    PRINCIPAL à qual está ligada (`NR_SECAO_PRINCIPAL` no cadastro de eleitorado);
  - a agregação muda de eleição para eleição (uma agregada vira principal, uma principal
    passa a ser agregada a outra);
  - seções novas surgem por divisão de seções cheias, e algumas trocam de escola.

Então "a seção X sumiu" quase sempre quer dizer "a seção X agora está agregada à Y" (ou
"a Y existia agregada à X em 2022"). Aqui as seções são ligadas em GRUPOS:

  1. toda seção agregada é ligada à sua principal, em cada ano;
  2. a mesma numeração (na zona) é a mesma seção nos dois anos;
  3. uma seção nova (sem número em 2022) é ligada às seções da mesma escola que já
     existiam — é a divisão de uma delas, e sem registro de qual foi, a escola inteira
     é a menor unidade honesta.

Os votos (só estão nas principais) são somados por grupo, o que torna 2022 e 2026
comparáveis. Cada grupo traz as seções principais de cada ano, escolas e bairro.

Uso:  python scripts/secoes_agregadas_municipio.py [CD_MUNICIPIO]   (padrão 11673, Piripiri)
"""

import os
import sys

import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import consolida_2022 as c22  # noqa: E402
import consolida_2026 as c26  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = os.path.join(RAIZ, 'public', 'dados') + os.sep
CD = int(sys.argv[1]) if len(sys.argv) > 1 else 11673
CACHE = os.path.join(RAIZ, 'dados', 'processados')

# ---- cadastro de eleitorado (seção -> principal / local / bairro) -------------------------
def cadastro(ano, carregador):
    cache = os.path.join(CACHE, f'_cadastro_{ano}_{CD}.pkl')
    if os.path.exists(cache):
        return pd.read_pickle(cache)
    e = carregador()
    e = e[e['CD_MUNICIPIO'] == str(CD)].copy()
    e.to_pickle(cache)
    return e


def normaliza(e, ano):
    d = e[['NR_ZONA', 'NR_SECAO', 'DS_TIPO_SECAO_AGREGADA', 'NR_SECAO_PRINCIPAL',
           'NR_LOCAL_VOTACAO', 'NM_LOCAL_VOTACAO', 'NM_BAIRRO', 'QT_ELEITOR_SECAO']].copy()
    d['zona'] = d['NR_ZONA'].astype(int)
    d['sec'] = d['NR_SECAO'].astype(int)
    d['papel'] = d['DS_TIPO_SECAO_AGREGADA'].map({'Principal': 'principal', 'Agregada': 'agregada'})
    assert d['papel'].notna().all(), f'{ano}: tipo de seção desconhecido'
    d['pai'] = d['NR_SECAO_PRINCIPAL'].astype(int).where(d['papel'] == 'agregada')
    d = d.rename(columns={'NR_LOCAL_VOTACAO': 'local', 'NM_LOCAL_VOTACAO': 'nm_local',
                          'NM_BAIRRO': 'bairro', 'QT_ELEITOR_SECAO': 'eleitores'})
    assert not d.duplicated(['zona', 'sec']).any(), f'{ano}: seção repetida no cadastro'
    return d[['zona', 'sec', 'papel', 'pai', 'local', 'nm_local', 'bairro', 'eleitores']]


a = normaliza(cadastro(2022, c22.carrega_eleitorado), 2022)
b = normaliza(cadastro(2026, c26.carrega_eleitorado), 2026)
cad = a.add_suffix('_22').rename(columns={'zona_22': 'zona', 'sec_22': 'sec'}).merge(
    b.add_suffix('_26').rename(columns={'zona_26': 'zona', 'sec_26': 'sec'}),
    on=['zona', 'sec'], how='outer')
cad['papel_22'] = cad['papel_22'].fillna('ausente')
cad['papel_26'] = cad['papel_26'].fillna('ausente')
cad['situacao'] = cad['papel_22'] + ' → ' + cad['papel_26']

# ---- grupos (union-find) ------------------------------------------------------------------
pai = {}


def raiz(x):
    while pai.setdefault(x, x) != x:
        pai[x] = pai[pai[x]]
        x = pai[x]
    return x


def une(x, y):
    pai[raiz(x)] = raiz(y)


for r in cad.itertuples():
    k = (r.zona, r.sec)
    raiz(k)
    for sufixo in ('22', '26'):
        p = getattr(r, f'pai_{sufixo}')
        if pd.notna(p):
            une(k, (r.zona, int(p)))

# seção nova (ou que sumiu) → junta com as seções da mesma escola presentes nos dois anos
for sufixo, outro in (('26', '22'), ('22', '26')):
    novas = cad[(cad[f'papel_{outro}'] == 'ausente') & (cad[f'papel_{sufixo}'] != 'ausente')]
    for r in novas.itertuples():
        local = getattr(r, f'local_{sufixo}')
        irmas = cad[(cad[f'local_{sufixo}'] == local) & (cad['zona'] == r.zona)
                    & (cad[f'papel_{outro}'] != 'ausente')]
        for s in irmas.itertuples():
            une((r.zona, r.sec), (s.zona, s.sec))

cad['grupo'] = [raiz((z, s)) for z, s in zip(cad['zona'], cad['sec'])]
ids = {g: i for i, g in enumerate(sorted(cad['grupo'].unique(), key=lambda g: (g[0], g[1])), start=1)}
cad['grupo'] = cad['grupo'].map(ids)

# ---- votos por seção principal (Presidente) ---------------------------------------------
con = duckdb.connect()
for t in ['fato_votos_secao', 'dim_local', 'dim_eleicao', 'dim_votavel']:
    con.sql(f"create view {t} as select * from '{D}{t}.parquet'")
sk = {an: con.sql(f'select SK_ELEICAO from dim_eleicao where ANO_ELEICAO={an} and FL_SERIE_PRINCIPAL').fetchone()[0]
      for an in (2022, 2026)}
NOMES = {(2022, 'LUIZ INÁCIO LULA DA SILVA'): 'lula', (2022, 'JAIR MESSIAS BOLSONARO'): 'rival',
         (2026, 'LUIZ INÁCIO LULA DA SILVA'): 'lula', (2026, 'FLAVIO NANTES BOLSONARO'): 'rival'}
v = con.sql(f"""
  with loc as (select distinct SK_LOCAL, SK_ELEICAO, NR_ZONA from dim_local where CD_MUNICIPIO={CD})
  select s.SK_ELEICAO, loc.NR_ZONA as zona, s.NR_SECAO as sec, d.NM_VOTAVEL, s.QT_VOTOS_NORM as v
  from fato_votos_secao s
  join loc on loc.SK_LOCAL = s.SK_LOCAL and loc.SK_ELEICAO = s.SK_ELEICAO
  join dim_votavel d on d.SK_VOTAVEL = s.SK_VOTAVEL
  where d.CD_CARGO = 1 and s.SK_ELEICAO in ({sk[2022]}, {sk[2026]})
""").df()
v['ano'] = v['SK_ELEICAO'].map({sk[2022]: 2022, sk[2026]: 2026})
v['papel'] = [NOMES.get((an, n), 'outros') for an, n in zip(v['ano'], v['NM_VOTAVEL'])]
v['zona'] = v['zona'].astype(int)
t = v.groupby(['zona', 'sec', 'ano', 'papel'])['v'].sum().unstack(['ano', 'papel'], fill_value=0)
t.columns = [f'{p}{str(an)[2:]}' for an, p in t.columns]
t = t.reset_index()
for c in ['lula22', 'rival22', 'outros22', 'lula26', 'rival26', 'outros26']:
    if c not in t:
        t[c] = 0
cad = cad.merge(t, on=['zona', 'sec'], how='left').fillna({c: 0 for c in t.columns if c not in ('zona', 'sec')})

# toda seção com voto tem de ser principal no cadastro do ano (agregada não vota)
assert (cad.loc[cad['papel_22'] != 'principal', ['lula22', 'rival22', 'outros22']].sum().sum() == 0), \
    'voto em seção que não é principal em 2022'
assert (cad.loc[cad['papel_26'] != 'principal', ['lula26', 'rival26', 'outros26']].sum().sum() == 0), \
    'voto em seção que não é principal em 2026'

# ---- por grupo ----------------------------------------------------------------------------
def junta(s):
    return ', '.join(str(int(x)) for x in sorted(s.dropna().unique()))


def nomes(s):
    return ' / '.join(sorted({str(x).title() for x in s.dropna()}))


linhas = []
for g, x in cad.groupby('grupo'):
    p22 = x[x['papel_22'] == 'principal']
    p26 = x[x['papel_26'] == 'principal']
    novas = x[(x['papel_22'] == 'ausente')]
    tipo = ('Mesmas seções principais' if set(p22['sec']) == set(p26['sec']) and len(novas) == 0
            else ('Com seção nova' if len(novas) else 'Agregação alterada'))
    r = {
        'grupo': g,
        'tipo': tipo,
        'bairro': nomes(x['bairro_26'].fillna(x['bairro_22'])),
        'escola': nomes(x['nm_local_26'].fillna(x['nm_local_22'])),
        'principais_2022': junta(p22['sec']),
        'principais_2026': junta(p26['sec']),
        'agregadas_2022': junta(x.loc[x['papel_22'] == 'agregada', 'sec']),
        'agregadas_2026': junta(x.loc[x['papel_26'] == 'agregada', 'sec']),
        'novas_2026': junta(novas['sec']),
        'eleitores22': x.loc[x['papel_22'] != 'ausente', 'eleitores_22'].sum(),
        'eleitores26': x.loc[x['papel_26'] != 'ausente', 'eleitores_26'].sum(),
    }
    for c in ['lula22', 'rival22', 'outros22', 'lula26', 'rival26', 'outros26']:
        r[c] = x[c].sum()
    linhas.append(r)
G = pd.DataFrame(linhas)
for an in ('22', '26'):
    G[f'validos{an}'] = G[f'lula{an}'] + G[f'rival{an}'] + G[f'outros{an}']
    for p in ('lula', 'rival'):
        G[f'{p}{an}_pct'] = (G[f'{p}{an}'] / G[f'validos{an}'] * 100).where(G[f'validos{an}'] > 0)
G['d_lula_votos'] = G['lula26'] - G['lula22']
G['d_lula_pp'] = G['lula26_pct'] - G['lula22_pct']
G['d_margem_pp'] = (G['lula26_pct'] - G['rival26_pct']) - (G['lula22_pct'] - G['rival22_pct'])
G = G.sort_values('d_lula_pp')

# conservação: nenhum voto pode se perder na agrupagem
assert G['lula22'].sum() == v[(v.ano == 2022) & (v.papel == 'lula')]['v'].sum()
assert G['lula26'].sum() == v[(v.ano == 2026) & (v.papel == 'lula')]['v'].sum()

# ---- seções "sem par" pela chave (escola, número): o que aconteceu com cada uma -------------
def explica(r):
    s = int(r.sec)
    if r.situacao == 'principal → agregada':
        return (f'Principal em 2022 (escola {r.nm_local_22.title()}); em 2026 foi agregada à seção '
                f'{int(r.pai_26)} ({r.nm_local_26.title()}) — seus votos agora entram na {int(r.pai_26)}.')
    if r.situacao == 'agregada → principal':
        return (f'Em 2022 era agregada à seção {int(r.pai_22)} ({r.nm_local_22.title()}), por isso '
                f'não aparece nos votos de 2022; em 2026 virou principal.')
    if r.situacao == 'ausente → principal':
        return f'Seção criada depois de 2022 (nº novo), em {r.nm_local_26.title()} — divisão de seção cheia.'
    if r.situacao == 'ausente → agregada':
        return (f'Seção criada depois de 2022 e já agregada à {int(r.pai_26)} '
                f'({r.nm_local_26.title()}).')
    if r.situacao == 'principal → principal' and r.local_22 != r.local_26:
        return (f'Principal nos dois anos, mas mudou de escola: {r.nm_local_22.title()} → '
                f'{r.nm_local_26.title()}.')
    if r.situacao == 'agregada → agregada' and r.pai_22 != r.pai_26:
        return f'Agregada nos dois anos, mas a principal mudou: {int(r.pai_22)} → {int(r.pai_26)}.'
    return ''


# Aviso oficial de mudança de local de funcionamento em 2026, colado pelo usuário em
# 11/10/2026. O TSE NÃO registra essas mudanças: cadastro e arquivo de votos mantêm o local
# cadastrado (Nenem Cavalcante, CHRISFAPI...). Fica aqui só como informação — o bairro e a
# escola das análises seguem o local cadastrado, que é onde estão os eleitores da seção.
AVISO_2026 = {
    16: 'Fundação Tia Zélia', 17: 'Fundação Tia Zélia', 18: 'Fundação Tia Zélia',
    95: 'Fundação Tia Zélia (agregada à 16)', 191: 'Fundação Tia Zélia (agregada à 17)',
    270: 'Fundação Tia Zélia (agregada à 18)', 238: 'Fundação Tia Zélia (removida da Secretaria de Agricultura)',
    49: 'Gerência Regional de Educação (GRE)', 50: 'Gerência Regional de Educação (GRE)',
    160: 'CAPS AD', 226: 'CAPS AD',
    60: 'Centro Educativo Antônio de Freitas Resende', 61: 'Centro Educativo Antônio de Freitas Resende',
    62: 'Centro Educativo Antônio de Freitas Resende',
    112: 'escola da localidade Pequi', 162: 'escola da localidade Pega Bem',
    180: 'casa onde está instalada a escola (Baixão)', 151: 'escola da localidade Santa Rosa',
    182: 'postinho da comunidade (Açude do Governo)', 183: 'Unidade Básica de Saúde (Oiticica II)',
    254: 'postinho ao lado da escola', 281: 'postinho ao lado da escola',
}
cad['funcionou_em_2026_aviso'] = cad['sec'].map(AVISO_2026)

cad['explicacao'] = [explica(r) for r in cad.itertuples()]
cad['mudou_escola'] = (cad['papel_22'] != 'ausente') & (cad['papel_26'] != 'ausente') & (cad['local_22'] != cad['local_26'])

pd.to_pickle({'cad': cad, 'G': G}, os.path.join(CACHE, f'_secoes_{CD}.pkl'))

# ---- Excel --------------------------------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = 'Leia-me'
txt = [
    'Seções de Piripiri entre 2022 e 2026 — o que aconteceu com cada uma',
    '',
    'Seção agregada não aparece nos votos: seus votos entram na seção principal. A agregação muda',
    'entre eleições, e seções novas surgem por divisão. Por isso "sumiu/apareceu" não é fim/criação',
    'de eleitorado: na maioria é a agregação mudando.',
    '',
    'Grupos: seções ligadas por agregação (em qualquer dos anos), mesma numeração na zona e, no caso de',
    'seção nova, as seções da mesma escola. Os votos são somados por grupo — é o nível em que',
    '2022 e 2026 são comparáveis. "Mesmas seções principais" = comparação seção a seção válida.',
    'Válidos = votos nominais em candidato a Presidente. p.p. = pontos percentuais.',
    '',
    'Aba "Aviso 2026": mudanças de local de funcionamento informadas oficialmente. O TSE não as',
    'registra (cadastro e votos mantêm o local cadastrado); as análises seguem o local cadastrado.',
]
for i, l in enumerate(txt, start=1):
    ws.cell(i, 1, l)
ws['A1'].font = Font(bold=True, size=13)
ws.column_dimensions['A'].width = 105

VERDE, VERMELHO = PatternFill('solid', fgColor='FFD9EAD3'), PatternFill('solid', fgColor='FFF4CCCC')


def aba(nome, df, cols, larguras):
    w = wb.create_sheet(nome)
    w.append([c[1] for c in cols])
    for c in w[1]:
        c.font = Font(bold=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor='FF1F2937')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for _, r in df.iterrows():
        w.append([None if pd.isna(r[k]) else (float(r[k]) if isinstance(r[k], float) else r[k]) for k, _, _ in cols])
    for j, (k, _, fmt) in enumerate(cols, start=1):
        w.column_dimensions[get_column_letter(j)].width = larguras.get(k, 14)
        for i in range(2, w.max_row + 1):
            c = w.cell(i, j)
            if fmt:
                c.number_format = fmt
            if k in ('d_lula_pp', 'd_lula_votos', 'd_margem_pp') and isinstance(c.value, (int, float)):
                c.fill = VERDE if c.value > 0 else (VERMELHO if c.value < 0 else PatternFill())
            if k in ('explicacao', 'escola'):
                c.alignment = Alignment(wrap_text=True, vertical='top')
    w.row_dimensions[1].height = 42
    w.freeze_panes = 'C2'
    w.auto_filter.ref = w.dimensions


INT, PCT, PP = '#,##0', '0.0', '+0.0;-0.0;0.0'
aba('Grupos', G, [
    ('grupo', 'Grupo', None), ('tipo', 'Tipo', None), ('bairro', 'Bairro', None), ('escola', 'Escola(s)', None),
    ('principais_2022', 'Principais 2022', None), ('principais_2026', 'Principais 2026', None),
    ('agregadas_2022', 'Agregadas 2022', None), ('agregadas_2026', 'Agregadas 2026', None),
    ('novas_2026', 'Seções novas em 2026', None),
    ('lula22', 'Lula 2022', INT), ('rival22', 'Bolsonaro 2022', INT), ('validos22', 'Válidos 2022', INT),
    ('lula22_pct', 'Lula 2022 %', PCT), ('lula26', 'Lula 2026', INT), ('rival26', 'Flávio 2026', INT),
    ('validos26', 'Válidos 2026', INT), ('lula26_pct', 'Lula 2026 %', PCT),
    ('d_lula_votos', 'Δ Lula (votos)', INT), ('d_lula_pp', 'Δ Lula (p.p.)', PP), ('d_margem_pp', 'Δ margem (p.p.)', PP),
], {'tipo': 24, 'bairro': 22, 'escola': 46, 'principais_2022': 20, 'principais_2026': 20,
    'agregadas_2022': 20, 'agregadas_2026': 20, 'novas_2026': 16})

mov = cad[cad['situacao'] != 'principal → principal'].copy()
mov = pd.concat([mov, cad[cad['mudou_escola']]]).drop_duplicates(['zona', 'sec']).sort_values('sec')
mov['bairro'] = mov['bairro_26'].fillna(mov['bairro_22']).str.title()
mov['escola_22'] = mov['nm_local_22'].str.title()
mov['escola_26'] = mov['nm_local_26'].str.title()
aba('Mudanças de seção', mov, [
    ('sec', 'Seção', None), ('situacao', '2022 → 2026', None), ('bairro', 'Bairro', None),
    ('escola_22', 'Escola 2022', None), ('escola_26', 'Escola 2026', None),
    ('pai_22', 'Principal 2022', None), ('pai_26', 'Principal 2026', None),
    ('eleitores_22', 'Eleitores 2022', INT), ('eleitores_26', 'Eleitores 2026', INT),
    ('grupo', 'Grupo', None), ('explicacao', 'O que aconteceu', None),
    ('funcionou_em_2026_aviso', 'Funcionou em 2026 (aviso oficial)', None),
], {'situacao': 24, 'bairro': 20, 'escola_22': 40, 'escola_26': 40, 'explicacao': 80,
    'funcionou_em_2026_aviso': 44})

# seções citadas no aviso, mesmo as que não mudaram de situação no cadastro
av = cad[cad['funcionou_em_2026_aviso'].notna()].copy().sort_values('sec')
av['bairro'] = av['bairro_26'].fillna(av['bairro_22']).str.title()
av['escola_26'] = av['nm_local_26'].str.title()
av['escola_22'] = av['nm_local_22'].str.title()
aba('Aviso 2026', av, [
    ('sec', 'Seção', None), ('situacao', '2022 → 2026', None), ('bairro', 'Bairro', None),
    ('escola_22', 'Escola 2022 (cadastro)', None), ('escola_26', 'Escola 2026 (cadastro e votos TSE)', None),
    ('funcionou_em_2026_aviso', 'Funcionou em 2026 (aviso oficial)', None), ('grupo', 'Grupo', None),
], {'situacao': 24, 'bairro': 20, 'escola_22': 40, 'escola_26': 42, 'funcionou_em_2026_aviso': 50})

saida = os.path.join(CACHE, 'secoes_agregadas_piripiri_2022_2026.xlsx')
wb.save(saida)
print('gravado:', saida)
print(cad['situacao'].value_counts().to_string())
print('grupos:', len(G), G['tipo'].value_counts().to_dict())
